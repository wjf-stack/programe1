from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT = r"d:\cursor\cursor_programe1\zhilianweiji\智联微记_网页插件主题_动态感汇报.pptx"

BLUE = RGBColor(47, 107, 255)
DEEP_BLUE = RGBColor(24, 70, 189)
LIGHT_BG = RGBColor(244, 247, 252)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(16, 32, 63)
MUTED = RGBColor(95, 111, 143)
LINE = RGBColor(219, 228, 245)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_BG
    bg.line.fill.background()

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.45))
    top.fill.solid()
    top.fill.fore_color.rgb = WHITE
    top.line.fill.background()

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(3.2), Inches(0.07))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()


def add_header(slide, left, right):
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(8), Inches(0.25))
    p = tx.text_frame.paragraphs[0]
    p.text = left
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DEEP_BLUE

    tx2 = slide.shapes.add_textbox(Inches(10.1), Inches(0.12), Inches(2.7), Inches(0.25))
    p2 = tx2.text_frame.paragraphs[0]
    p2.text = right
    p2.alignment = PP_ALIGN.RIGHT
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED


def add_title(slide, title, subtitle=""):
    t = slide.shapes.add_textbox(Inches(0.75), Inches(0.9), Inches(8.5), Inches(1.4))
    p = t.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEXT
    if subtitle:
        p2 = t.text_frame.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = MUTED


def add_card(slide, x, y, w, h, title, body_lines):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE

    tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.16), Inches(w - 0.4), Inches(h - 0.3))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = DEEP_BLUE

    for line in body_lines:
        px = tf.add_paragraph()
        px.text = f"• {line}"
        px.font.size = Pt(13)
        px.font.color.rgb = TEXT


# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_header(s, "智联微记｜网页插件主题汇报", "简约科技风 / 正式展示")
add_title(s, "让网页信息\n快速沉淀为可复用知识", "浏览器插件 + Web 应用 + 知识图谱 的一体化方案")

hero = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.0), Inches(4.6), Inches(5.7))
hero.fill.solid(); hero.fill.fore_color.rgb = WHITE; hero.line.color.rgb = LINE

for i, txt in enumerate(["开始截图摘录", "打开知识图谱", "打开仪表盘"]):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(2.0 + i * 1.05), Inches(3.9), Inches(0.78))
    b.fill.solid(); b.fill.fore_color.rgb = BLUE if i != 1 else RGBColor(77, 125, 255)
    b.line.fill.background()
    pt = b.text_frame.paragraphs[0]
    pt.text = txt
    pt.alignment = PP_ALIGN.CENTER
    pt.font.bold = True
    pt.font.size = Pt(18)
    pt.font.color.rgb = WHITE

for i, (n, l) in enumerate([("3端", "插件/前端/后端"), ("4步", "采集→整理→关联→复用"), ("3层", "主题→分支→知识点")]):
    add_card(s, 0.9 + i * 2.35, 5.25, 2.15, 1.2, n, [l])

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_header(s, "核心功能矩阵", "Core Features")
add_title(s, "网页插件能力与业务价值")
add_card(s, 0.8, 2.0, 6.0, 4.8, "插件侧能力", [
    "任意网页拖拽选区截图摘录",
    "主题 + 分支结构化录入",
    "保存中/成功/失败状态反馈",
    "一键跳转知识图谱与仪表盘"
])
add_card(s, 6.9, 2.0, 5.6, 4.8, "Web 端能力", [
    "仪表盘管理：查看、编辑、删除",
    "笔记编辑增强：图片粘贴上传",
    "3D知识图谱：多星系分簇展示",
    "主题/分支筛选与同分支聚焦"
])

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_header(s, "操作流程演示", "User Journey")
add_title(s, "从“看到信息”到“知识沉淀”")

steps = [
    ("01 启动插件", "点击“开始截图摘录”进入选区模式"),
    ("02 结构化录入", "填写主题与分支，保存截图笔记"),
    ("03 仪表盘管理", "统一维护笔记、标签与内容"),
    ("04 图谱复盘", "按主题/分支聚焦知识关系"),
]
for i, (t, d) in enumerate(steps):
    add_card(s, 0.8 + i * 3.15, 2.2, 2.9, 2.0, t, [d])

add_card(s, 0.8, 4.5, 4.0, 1.9, "截图占位 1", ["插件弹窗截图"])
add_card(s, 4.95, 4.5, 4.0, 1.9, "截图占位 2", ["摘录浮层截图"])
add_card(s, 9.1, 4.5, 3.4, 1.9, "截图占位 3", ["图谱筛选截图"])

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_header(s, "技术架构", "Architecture")
add_title(s, "插件 + 前后端 + 图谱引擎协同")
add_card(s, 0.8, 2.0, 6.0, 4.8, "技术栈", [
    "前端：Vue3 + Vite + Pinia + Router",
    "后端：Node.js + Express + LowDB + JWT",
    "扩展：Manifest V3 + Content/Popup/SW",
    "可视化：Three.js + GLTFLoader"
])
add_card(s, 6.9, 2.0, 5.6, 4.8, "关键亮点", [
    "跨端闭环：采集到复盘一体化",
    "结构化维度：主题 + 分支 + 标签",
    "图谱层级：多星系 + 分支轨道",
    "可维护：模块清晰、便于扩展"
])

# Slide 5
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
add_header(s, "价值与总结", "Conclusion")
add_title(s, "美观 + 实用，适配正式展示")
add_card(s, 0.8, 2.0, 6.0, 4.8, "应用价值", [
    "学习场景：高频信息快速沉淀",
    "课程汇报：插件演示更直观",
    "产品介绍：问题-方案-效果链路完整"
])
add_card(s, 6.9, 2.0, 5.6, 4.8, "后续规划", [
    "语义推荐与关联挖掘",
    "多人协作与权限控制",
    "多端同步与稳定性提升"
])

note = s.shapes.add_textbox(Inches(0.9), Inches(6.95), Inches(12.0), Inches(0.35))
p = note.text_frame.paragraphs[0]
p.text = "演示建议：放映模式下配合“切换（淡化）”转场；第3页替换真实截图即可用于答辩。"
p.font.size = Pt(12)
p.font.color.rgb = MUTED

prs.save(OUTPUT)
print(OUTPUT)
