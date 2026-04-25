from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = r"d:\cursor\cursor_programe1\zhilianweiji\智联微记_网页插件主题_精美增强版.pptx"

C_BG = RGBColor(242, 246, 255)
C_WHITE = RGBColor(255, 255, 255)
C_BLUE = RGBColor(47, 107, 255)
C_BLUE_D = RGBColor(23, 63, 160)
C_TXT = RGBColor(18, 36, 74)
C_MUTED = RGBColor(93, 112, 150)
C_LINE = RGBColor(214, 226, 249)
C_ACC = RGBColor(98, 144, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bg(slide):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = C_BG
    r.line.fill.background()
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.42))
    top.fill.solid(); top.fill.fore_color.rgb = C_WHITE
    top.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(3.4), Inches(0.07))
    accent.fill.solid(); accent.fill.fore_color.rgb = C_BLUE
    accent.line.fill.background()


def header(slide, l, r):
    t = slide.shapes.add_textbox(Inches(0.46), Inches(0.10), Inches(8.8), Inches(0.25))
    p = t.text_frame.paragraphs[0]
    p.text = l
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_BLUE_D

    t2 = slide.shapes.add_textbox(Inches(9.6), Inches(0.10), Inches(3.1), Inches(0.25))
    p2 = t2.text_frame.paragraphs[0]
    p2.text = r
    p2.alignment = PP_ALIGN.RIGHT
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_MUTED


def title(slide, main, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.72), Inches(0.72), Inches(11.5), Inches(1.1))
    p = tb.text_frame.paragraphs[0]
    p.text = main
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = C_TXT
    if sub:
        p2 = tb.text_frame.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(14)
        p2.font.color.rgb = C_MUTED


def card(slide, x, y, w, h, htxt, lines):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid(); c.fill.fore_color.rgb = C_WHITE
    c.line.color.rgb = C_LINE

    tb = slide.shapes.add_textbox(Inches(x+0.20), Inches(y+0.14), Inches(w-0.35), Inches(h-0.25))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = htxt
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_BLUE_D
    for ln in lines:
        px = tf.add_paragraph()
        px.text = f"• {ln}"
        px.font.size = Pt(12.2)
        px.font.color.rgb = C_TXT


def kpi(slide, x, y, n, lab):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.35), Inches(1.1))
    c.fill.solid(); c.fill.fore_color.rgb = C_WHITE
    c.line.color.rgb = C_LINE
    tb = slide.shapes.add_textbox(Inches(x+0.14), Inches(y+0.10), Inches(2.08), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = n
    p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = C_BLUE_D
    p2 = tb.text_frame.add_paragraph()
    p2.text = lab
    p2.font.size = Pt(11); p2.font.color.rgb = C_MUTED


def step_box(slide, x, y, t, d):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.03), Inches(1.55))
    c.fill.solid(); c.fill.fore_color.rgb = C_WHITE
    c.line.color.rgb = C_LINE
    tb = slide.shapes.add_textbox(Inches(x+0.15), Inches(y+0.11), Inches(2.7), Inches(1.3))
    p = tb.text_frame.paragraphs[0]
    p.text = t
    p.font.size = Pt(14.5); p.font.bold = True; p.font.color.rgb = C_BLUE_D
    p2 = tb.text_frame.add_paragraph()
    p2.text = d
    p2.font.size = Pt(11.2); p2.font.color.rgb = C_MUTED


def bar(slide, x, y, label, v, maxv=100):
    t = slide.shapes.add_textbox(Inches(x), Inches(y-0.02), Inches(2.2), Inches(0.2))
    p = t.text_frame.paragraphs[0]
    p.text = label
    p.font.size = Pt(11)
    p.font.color.rgb = C_MUTED

    bgbar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+1.6), Inches(y), Inches(3.1), Inches(0.22))
    bgbar.fill.solid(); bgbar.fill.fore_color.rgb = RGBColor(227, 236, 252); bgbar.line.fill.background()
    w = 3.1 * (v/maxv)
    fg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+1.6), Inches(y), Inches(w), Inches(0.22))
    fg.fill.solid(); fg.fill.fore_color.rgb = C_ACC; fg.line.fill.background()

    tv = slide.shapes.add_textbox(Inches(x+4.8), Inches(y-0.03), Inches(0.7), Inches(0.22))
    pv = tv.text_frame.paragraphs[0]
    pv.text = f"{v}%"
    pv.font.size = Pt(11); pv.font.bold = True; pv.font.color.rgb = C_BLUE_D


# 1 封面
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "智联微记｜网页插件主题精美汇报", "课程展示 / 产品介绍")
title(s, "网页信息，一键沉淀为结构化知识资产", "以浏览器插件为入口，连接仪表盘管理与3D知识图谱探索")

mock = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.3), Inches(4.5), Inches(5.6))
mock.fill.solid(); mock.fill.fore_color.rgb = C_WHITE; mock.line.color.rgb = C_LINE
for i, txt in enumerate(["开始截图摘录", "打开知识图谱", "打开仪表盘", "退出登录"]):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(2.1+i*0.95), Inches(3.8), Inches(0.72))
    b.fill.solid(); b.fill.fore_color.rgb = C_BLUE if i < 3 else RGBColor(240, 244, 253)
    b.line.color.rgb = C_LINE
    p = b.text_frame.paragraphs[0]; p.text = txt; p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(17); p.font.bold = True; p.font.color.rgb = C_WHITE if i < 3 else C_BLUE_D

kpi(s, 0.9, 5.35, "3端协同", "插件/前端/后端")
kpi(s, 3.45, 5.35, "4步闭环", "采集→整理→关联→复用")
kpi(s, 6.0, 5.35, "3层图谱", "主题→分支→知识点")

# 2 项目背景
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "项目背景与目标", "Why This Product")
title(s, "痛点驱动：碎片化信息难沉淀、难复盘")
card(s, 0.8, 2.0, 6.0, 4.9, "现有问题", [
    "网页学习信息分散，截图与笔记割裂",
    "收藏夹堆积后难检索、难形成知识结构",
    "笔记缺乏主题与分支维度，复盘效率低",
    "知识点之间关联弱，无法形成可视网络"
])
card(s, 6.95, 2.0, 5.55, 4.9, "项目目标", [
    "构建“插件采集 + 系统管理 + 图谱复盘”闭环",
    "在采集阶段完成主题与分支结构化",
    "通过图谱展示关系，提升学习迁移能力",
    "提供可用于课程答辩的完整演示方案"
])

# 3 功能全景
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "功能全景", "Feature Map")
title(s, "核心能力模块化设计")
card(s, 0.8, 2.0, 4.05, 2.15, "A. 插件采集", ["框选截图", "主题/分支必填", "保存状态反馈"])
card(s, 5.0, 2.0, 4.05, 2.15, "B. 仪表盘管理", ["笔记列表与检索", "删除与编辑", "主题管理"])
card(s, 9.2, 2.0, 3.2, 2.15, "C. 图谱探索", ["多星系分簇", "主题/分支筛选"])
card(s, 0.8, 4.35, 6.45, 2.35, "D. 编辑增强", ["图片粘贴上传", "标签维护", "快捷保存"]) 
card(s, 7.4, 4.35, 5.1, 2.35, "E. 系统能力", ["登录鉴权(JWT)", "文件上传(Multer)", "轻量存储(LowDB)"]) 

# 4 用户流程
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "用户流程", "User Journey")
title(s, "操作路径：从采集到复盘")
steps = [
    ("01 打开插件", "登录后点击“开始截图摘录”，进入选区模式"),
    ("02 结构化录入", "输入标题、主题、分支；保存截图与来源"),
    ("03 仪表盘整理", "查看、编辑、删除笔记；管理主题标签"),
    ("04 图谱复盘", "按主题/分支聚焦，查看关联与节点详情")
]
for idx, st in enumerate(steps):
    step_box(s, 0.8 + idx*3.15, 2.15, st[0], st[1])

card(s, 0.8, 4.25, 3.9, 2.15, "截图位1", ["插件弹窗", "建议使用真实界面图"])
card(s, 4.95, 4.25, 3.9, 2.15, "截图位2", ["摘录浮层", "展示主题分支输入"])
card(s, 9.1, 4.25, 3.2, 2.15, "截图位3", ["图谱页面", "展示筛选与高亮"])

# 5 插件页细化
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "插件能力细化", "Plugin UX")
title(s, "插件体验设计：轻操作 + 强反馈")
card(s, 0.8, 2.0, 6.1, 4.9, "交互细节", [
    "按钮分层：采集、图谱、仪表盘、登录状态",
    "摘录入口显著，避免学习流程中断",
    "错误提示可读：未登录/参数缺失/端口异常",
    "支持端口适配，提升本地开发可用性"
])
card(s, 7.05, 2.0, 5.45, 4.9, "工程实现", [
    "Content Script 负责页面交互与选区",
    "Service Worker 负责消息与上传请求",
    "Popup 负责登录状态与快捷操作",
    "统一消息协议：capture/save/search/open"
])

# 6 图谱页细化
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "图谱设计", "Knowledge Graph")
title(s, "3D图谱：从“列表管理”走向“关系认知")
card(s, 0.8, 2.0, 6.1, 4.9, "视觉结构", [
    "星系：按关联连通组件分簇",
    "恒星：主题（Tag）节点中心",
    "轨道：分支（Branch）层级表达",
    "行星：具体知识点笔记"
])
card(s, 7.05, 2.0, 5.45, 4.9, "交互能力", [
    "主题筛选 + 分支筛选双维过滤",
    "点击节点查看标题/内容/图片预览",
    "同分支聚焦，一键查看分支全量节点",
    "摄像机重置、双击聚焦、流星动态效果"
])

# 7 技术架构
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "技术架构", "Architecture")
title(s, "模块分层清晰，便于维护与扩展")
card(s, 0.8, 2.0, 6.1, 2.3, "前端层", ["Vue3 + Vite + Pinia + Vue Router", "页面：Dashboard / NoteEditor / Graph / Login"])
card(s, 0.8, 4.45, 6.1, 2.2, "后端层", ["Express REST API", "JWT鉴权 + Multer上传 + LowDB存储"])
card(s, 7.05, 2.0, 5.45, 4.65, "扩展层", [
    "Manifest V3（Popup / Content / Service Worker）",
    "采集消息流：选区→截图→上传→笔记创建",
    "跳转消息流：插件→图谱/仪表盘页面",
    "接口协同：/auth /notes /tags /graph /upload"
])

# 8 数据模型
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "数据模型", "Data Schema")
title(s, "结构化字段支撑知识复用")
card(s, 0.8, 2.0, 4.2, 4.9, "核心实体", [
    "User：账号、密码哈希、创建时间",
    "Note：标题、内容、类型、来源、branch",
    "Tag：主题标签",
    "Relation：节点关系"
])
card(s, 5.2, 2.0, 3.9, 4.9, "关联关系", [
    "noteTags：笔记与主题多对多",
    "Graph Relations：sourceId → targetId",
    "branch字段支持分支轨道展示"
])
card(s, 9.3, 2.0, 3.2, 4.9, "兼容策略", [
    "新数据优先读取branch字段",
    "旧数据可从正文提取“分支：xxx”",
    "保证历史内容可视化不失效"
])

# 9 难点与解决
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "难点与解决", "Engineering Challenges")
title(s, "关键难点与落地方案")
card(s, 0.8, 2.0, 6.0, 4.9, "难点", [
    "扩展采集交互复杂：选区、截图、消息通信",
    "本地端口不稳定：5173/5174/5175动态变化",
    "图谱节点易拥挤：单中心布局可读性差",
    "结构化演进：从文本分支到字段分支"
])
card(s, 6.95, 2.0, 5.55, 4.9, "解决方案", [
    "分层脚本设计：Popup/Content/SW职责清晰",
    "端口探测与容错提示，提升可用性",
    "多星系分簇 + 分支轨道，提升可读性",
    "后端新增branch并做旧数据兼容解析"
])

# 10 成果指标
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "成果与评估", "Outcomes")
title(s, "可视化展示 + 工程可用性双达成")
card(s, 0.8, 1.95, 5.8, 2.1, "阶段成果", [
    "完成插件采集、仪表盘、图谱三端联调",
    "支持主题/分支筛选与同分支聚焦",
    "支持笔记编辑、删除、图片粘贴上传"
])
card(s, 0.8, 4.2, 5.8, 2.45, "展示价值", [
    "适用于课程汇报、毕业设计答辩、产品Demo",
    "叙事完整：背景→方案→实现→效果→规划"
])

card(s, 6.85, 1.95, 5.65, 4.7, "效果评分（建议展示项）", ["以下可在现场按真实数据更新"]) 
bar(s, 7.1, 2.95, "功能完整度", 92)
bar(s, 7.1, 3.45, "交互体验", 88)
bar(s, 7.1, 3.95, "视觉表现", 90)
bar(s, 7.1, 4.45, "工程稳定性", 85)
bar(s, 7.1, 4.95, "可扩展性", 87)

# 11 结束页
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s); header(s, "总结与规划", "Conclusion")
title(s, "智联微记：把碎片信息变成知识网络", "感谢聆听 · Q&A")
card(s, 0.8, 2.0, 6.0, 4.9, "本项目关键词", [
    "网页插件采集", "结构化沉淀", "图谱化复盘", "教学与展示友好"
])
card(s, 6.95, 2.0, 5.55, 4.9, "下一步计划", [
    "语义检索与相似知识推荐",
    "团队协作与知识共享权限",
    "多端同步与离线数据策略",
    "模型辅助总结与自动归档"
])

prs.save(OUT)
print(OUT)
